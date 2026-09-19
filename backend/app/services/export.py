"""Exportación de resultados: JSON completo, CSV y Excel. Opción `redact` para versiones sin PII."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

import numpy as np
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (AudioSegment, Call, CriticalEvent, EmotionPrediction, MLModel, SatisfactionScore, Speaker,
                        Transcription)
from app.services import narrative
from app.services import serializers as ser

EMOTIONS = ["angry", "disgust", "fear", "happy", "neutral", "sad", "surprise"]


def gather(db: Session, call: Call, redact: bool = False, roles_labels: dict | None = None) -> dict:
    """Todos los datos de una llamada en estructuras planas (base de JSON/CSV/Excel/PDF)."""
    lm = {s.id: s.label for s in call.speakers}
    segs = db.scalars(select(AudioSegment).where(AudioSegment.call_id == call.id).order_by(AudioSegment.start)).all()
    trs = db.scalars(select(Transcription).where(Transcription.call_id == call.id).order_by(Transcription.start)).all()
    emos = db.scalars(select(EmotionPrediction).where(EmotionPrediction.call_id == call.id).order_by(EmotionPrediction.start)).all()
    sats = db.scalars(select(SatisfactionScore).where(SatisfactionScore.call_id == call.id)).all()
    evs = db.scalars(select(CriticalEvent).where(CriticalEvent.call_id == call.id).order_by(CriticalEvent.timestamp)).all()
    model = db.get(MLModel, call.model_id) if call.model_id else None
    sat = {"speakers": {}, "interaction": None}
    for r in sats:
        d = {"score": r.score, "confidence": r.confidence, "trend": r.trend, "initial_score": r.initial_score,
             "final_score": r.final_score, "engine_version": r.engine_version, **(r.details or {})}
        if r.scope == "interaction":
            sat["interaction"] = d
        else:
            sat["speakers"][lm[r.speaker_id]] = d
    return {
        "call": ser.call_detail(call, roles_labels),
        "speakers": [ser.speaker_out(s, roles_labels) for s in call.speakers],
        "segments": [{"speaker": lm[s.speaker_id], "start": s.start, "end": s.end, "overlap": s.overlap, "source": s.source} for s in segs],
        "transcription": [ser.transcription_out(t, redact, lm[t.speaker_id]) for t in trs],
        "emotions": [ser.emotion_out(p, lm[p.speaker_id]) for p in emos],
        "satisfaction": sat,
        "events": [ser.event_out(e, lm.get(e.speaker_id)) for e in evs],
        "model": None if model is None else {"id": model.id, "name": model.name, "version": model.version, "kind": model.kind,
                                             "language": model.language, "base_model": model.base_model, "hf_id": model.hf_id},
        "confidence": {"analysis_quality": call.analysis_quality, "audio_quality": call.audio_quality,
                       "speaker_confidence": {k: v["confidence"] for k, v in sat["speakers"].items()}},
        "interaction": call.interaction,
        "config": call.config_snapshot,
        "redacted": redact,
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "notice": "Estimaciones de un modelo; no constituyen afirmaciones definitivas. Separe los datos observados "
                  "(probabilidades, tiempos, transcripción) de la interpretación (satisfacción, eventos).",
    }


def _sat_at(timeline: list[dict], t: float) -> float | None:
    if not timeline:
        return None
    return round(float(np.interp(t, [p["t"] for p in timeline], [p["score"] for p in timeline])), 1)


def to_json(data: dict) -> bytes:
    import json
    d = dict(data)
    if d.get("config"):
        d["config"] = d["config"]                     # configuración usada (trazabilidad)
    return json.dumps(d, ensure_ascii=False, indent=2, default=str).encode("utf-8")


CSV_COLUMNS = ["call_id", "speaker", "start", "end", *EMOTIONS, "emotion", "confidence", "satisfaction"]


def csv_rows(data: dict) -> list[list]:
    cid = data["call"]["id"]
    out = []
    for p in data["emotions"]:
        tl = (data["satisfaction"]["speakers"].get(p["speaker"]) or {}).get("timeline", [])
        pr = p["probabilities"]
        out.append([cid, p["speaker"], round(p["start"], 2), round(p["end"], 2)] + [round(pr.get(e, 0.0), 5) for e in EMOTIONS]
                   + [p["emotion"], round(p["confidence"], 4), _sat_at(tl, (p["start"] + p["end"]) / 2)])
    return out


def to_csv(data: dict) -> bytes:
    """Columnas: call_id, speaker, start, end, <7 emociones>, emotion, confidence, satisfaction."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CSV_COLUMNS)
    w.writerows(csv_rows(data))
    return ("﻿" + buf.getvalue()).encode("utf-8")           # BOM: Excel abre bien los acentos


def to_xlsx(data: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")

    def sheet(name, header, rows, widths=None):
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, wd in enumerate(widths or [], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = wd
        return ws

    ws = wb.active
    ws.title = "Resumen"
    c = data["call"]
    ws.append(["Concepto", "Valor"])
    for cell in ws[1]:
        cell.font, cell.fill = hdr_font, hdr_fill
    rows = [("ID", c["id"]), ("Archivo", c["filename"]), ("Duración (s)", c["duration"]), ("Idioma", c["language"]),
            ("Estado", c["status"]), ("Modelo", c["model_version"]), ("Calidad de audio", c["audio_quality"]),
            ("Calidad del análisis", c["analysis_quality"]), ("Exportado", data["exported_at"])]
    for r in rows:
        ws.append(list(r))
    ws.append([])
    ws.append(["Participante", "Satisfacción", "Confianza", "Predominante", "Inicial", "Final", "Frustración %", "Tensión %",
               "Positivas %", "Negativas %", "Estabilidad %"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = hdr_font, hdr_fill
    names = {s["label"]: s["name"] for s in data["speakers"]}
    for lab, s in (c.get("summary") or {}).get("speakers", {}).items():
        m = s["metrics"]
        ws.append([names.get(lab, lab), s["score"], round(s["confidence"], 3), m["dominant_emotion"], m["initial_emotion"],
                   m["final_emotion"], round(m["frustration"] * 100, 1), round(m["tension"] * 100, 1),
                   round(m["positive"] * 100, 1), round(m["negative"] * 100, 1), round(m["stability"] * 100, 1)])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30
    sheet("Emociones", CSV_COLUMNS, csv_rows(data), [34, 12, 9, 9] + [9] * 7 + [11, 11, 13])
    sheet("Transcripción", ["speaker", "start", "end", "text", "confidence"],
          [[t["speaker"], round(t["start"], 2), round(t["end"], 2), t["text"], t["confidence"]] for t in data["transcription"]],
          [12, 9, 9, 100, 11])
    sheet("Eventos", ["timestamp", "speaker", "event", "emotion", "confidence", "satisfaction", "description"],
          [[e["timestamp"], e["speaker"], e["label"], e["emotion"], e["confidence"], e["satisfaction"], e["description"]]
           for e in data["events"]], [11, 12, 30, 10, 11, 12, 100])
    tl_rows = []
    for lab, s in data["satisfaction"]["speakers"].items():
        tl_rows += [[lab, p["t"], p["score"]] for p in s.get("timeline", [])]
    sheet("Satisfacción", ["speaker", "t", "score"], tl_rows, [14, 10, 10])
    for w in wb.worksheets:
        for row in w.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
