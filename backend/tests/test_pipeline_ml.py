"""Pruebas del pipeline con modelos REALES (Whisper + wav2vec baseline). Ejecutar con RUN_ML_TESTS=1.

La llamada de ejemplo es voz sintética (TTS): valida la mecánica de extremo a extremo (audio, diarización,
transcripción, PII, emociones, satisfacción, eventos, exportación), NO la calidad del reconocimiento emocional.
"""
import io
import json
import time
import zipfile

import pytest

from tests.conftest import EXAMPLES, make_tiny_emotion_model

pytestmark = [pytest.mark.ml, pytest.mark.skipif(not (EXAMPLES / "llamada_ejemplo_estereo.wav").exists(), reason="falta la llamada de ejemplo")]
API = "/api/v1"


def wait_completed(client, headers, cid, timeout=900):
    t0 = time.time()
    while time.time() - t0 < timeout:
        st = client.get(f"{API}/calls/{cid}/status", headers=headers).json()
        if st["status"] in ("COMPLETED", "ERROR"):
            return st
        time.sleep(3)
    raise TimeoutError("la llamada no terminó a tiempo")


@pytest.fixture(scope="module")
def completed_call(client, admin_a):
    with open(EXAMPLES / "llamada_ejemplo_estereo.wav", "rb") as f:
        r = client.post(f"{API}/calls/upload", headers=admin_a, files={"file": ("llamada_ejemplo_estereo.wav", f)})
    assert r.status_code == 201, r.text
    cid = r.json()["id"]
    st = wait_completed(client, admin_a, cid)
    assert st["status"] == "COMPLETED", st
    return cid


def test_end_to_end_results(client, admin_a, completed_call):
    cid = completed_call
    call = client.get(f"{API}/calls/{cid}", headers=admin_a).json()
    assert call["diarization_mode"] == "channels" and call["language"] == "es" and len(call["speakers"]) == 2
    assert call["allow_training"] is False and call["model_version"].startswith("emotion_baseline_wav2vec")
    assert call["audio_quality"] is not None and call["analysis_quality"] is not None
    assert any("entrenado para 'en'" in w for w in call["warnings"])              # aviso de idioma del baseline

    tr = client.get(f"{API}/calls/{cid}/transcription", headers=admin_a).json()["items"]
    text = " ".join(t["text"] for t in tr).lower()
    assert "molesto" in text and "gracias" in text
    assert {t["speaker"] for t in tr} == {"SPEAKER_00", "SPEAKER_01"}
    assert all({"start", "end", "confidence"} <= set(t) and t["end"] > t["start"] for t in tr)
    red = client.get(f"{API}/calls/{cid}/transcription", headers=admin_a, params={"redact": True}).json()["items"]
    joined = " ".join(t["text"] for t in red)
    assert "[PERSONA]" in joined and "[TELEFONO]" in joined and "Juan" not in joined

    emo = client.get(f"{API}/calls/{cid}/emotions", headers=admin_a).json()["items"]
    assert len(emo) > 10
    for e in emo:
        assert set(e["probabilities"]) == {"angry", "disgust", "fear", "happy", "neutral", "sad", "surprise"}
        assert abs(sum(e["probabilities"].values()) - 1) < 1e-3 and e["model"] and e["model_version"] and e["duration"] > 0
    assert emo[0]["start"] <= emo[1]["start"]
    # llamada en español + modelo de audio en inglés: se aplana el audio y se fusiona con la emoción del TEXTO (nativa)
    src = emo[0]["sources"]
    assert src["weights"]["text"] > src["weights"]["audio"] and src["temperature"] == 7.5
    with_text = [e for e in emo if e["sources"]["text"]]
    assert len(with_text) >= 0.6 * len(emo) and all(0 <= e["sources"]["agreement"] <= 1 for e in with_text)
    cust = [e for e in emo if e["speaker"] == "SPEAKER_00"]                       # el cliente: se queja al inicio y agradece al final
    assert any(e["emotion"] == "angry" for e in cust[:4]) and any(e["emotion"] == "happy" for e in cust[-3:])

    sat = client.get(f"{API}/calls/{cid}/satisfaction", headers=admin_a).json()
    assert set(sat["speakers"]) == {"SPEAKER_00", "SPEAKER_01"} and sat["interaction"]
    for s in sat["speakers"].values():
        assert 0 <= s["score"] <= 100 and 0 <= s["confidence"] <= 1 and s["timeline"] and s["factors"] is not None
    c = sat["speakers"]["SPEAKER_00"]
    assert c["final_score"] > c["initial_score"] + 10 and c["metrics"]["text_coverage"] > 0.5      # el guion es «molesto → satisfecho»
    ev = client.get(f"{API}/calls/{cid}/events", headers=admin_a).json()
    assert all({"timestamp", "event_type", "confidence", "label"} <= set(e) for e in ev)
    assert call["interaction"]["talk_time"]["SPEAKER_00"] > 5


def test_exports_and_pdf(client, admin_a, completed_call):
    cid = completed_call
    csv = client.get(f"{API}/calls/{cid}/export", headers=admin_a, params={"format": "csv"})
    header = csv.content.decode("utf-8-sig").splitlines()[0].split(",")
    assert header[:11] == ["call_id", "speaker", "start", "end", "angry", "disgust", "fear", "happy", "neutral", "sad", "surprise"]
    assert "satisfaction" in header
    js = client.get(f"{API}/calls/{cid}/export", headers=admin_a, params={"format": "json"}).json()
    assert {"call", "speakers", "segments", "transcription", "emotions", "satisfaction", "events", "model", "confidence"} <= set(js)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(client.get(f"{API}/calls/{cid}/export", headers=admin_a, params={"format": "xlsx"}).content))
    assert {"Resumen", "Emociones", "Transcripción", "Eventos"} <= set(wb.sheetnames)
    pdf = client.get(f"{API}/reports/{cid}", headers=admin_a)
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF" and len(pdf.content) > 20_000
    red = client.get(f"{API}/calls/{cid}/export", headers=admin_a, params={"format": "json", "redact": True}).json()
    assert "Juan" not in json.dumps(red["transcription"], ensure_ascii=False)


def test_audio_playback_range(client, admin_a, completed_call):
    r = client.get(f"{API}/calls/{completed_call}/audio", headers={**admin_a, "Range": "bytes=100-299"})
    assert r.status_code == 206 and r.headers["content-type"] == "audio/mpeg" and len(r.content) == 200


def test_mono_diarization_fallback_vs_ground_truth(tmp_path):
    """Mide la calidad REAL del diarizador de respaldo (sin pyannote) sobre el audio mono con verdad conocida."""
    import numpy as np
    from app.core.config import get_config_store
    from app.ml.audio.processing import load_vad, process_audio
    from app.ml.diarization.embedding import EmbeddingDiarizer
    from app.ml.diarization.spectral import SpectralDiarizer
    cfg = get_config_store().defaults("audio")
    res = process_audio(EXAMPLES / "llamada_ejemplo_mono.wav", tmp_path, cfg)
    assert res["mode"] == "mono"
    d = EmbeddingDiarizer().diarize(tmp_path, res["files"], load_vad(tmp_path), cfg)            # respaldo por defecto (embeddings)
    d_mfcc = SpectralDiarizer().diarize(tmp_path, res["files"], load_vad(tmp_path), cfg, hint=2)  # último recurso, forzado a 2
    truth = json.loads((EXAMPLES / "llamada_ejemplo_verdad.json").read_text(encoding="utf-8"))["turns"]
    grid = np.arange(0, res["duration"], 0.1)

    def lab(turns, key):
        out = np.full(len(grid), -1)
        for t in turns:
            s, e, sp = (t["start"], t["end"], t["speaker"]) if isinstance(t, dict) else (t.start, t.end, t.speaker)
            out[(grid >= s) & (grid < e)] = int(sp[-1])
        return out
    gt, pr = lab(truth, "t"), lab(d.turns, "p")
    m = (gt >= 0) & (pr >= 0)
    acc = max((gt[m] == pr[m]).mean(), (gt[m] != pr[m]).mean())                 # mejor asignación de etiquetas
    gt2, pr2 = lab(truth, "t"), lab(d_mfcc.turns, "p")
    m2 = (gt2 >= 0) & (pr2 >= 0)
    acc2 = max((gt2[m2] == pr2[m2]).mean(), (gt2[m2] != pr2[m2]).mean())
    print(f"diarización de respaldo (embeddings): acierto por trama = {acc:.2%}, calidad = {d.quality} | MFCC forzado: {acc2:.2%}")
    assert d.quality <= 70 and d.warnings                                       # nunca se presenta como diarización de alta calidad
    assert d.n_speakers_detected == 2 and acc > 0.85 and acc2 > 0.7


def test_resume_from_checkpoint_after_failure(client, monkeypatch):
    """Si falla la etapa de emociones a mitad, al reanudar NO se repiten audio/diarización/transcripción ni se duplican predicciones."""
    from sqlalchemy import func, select
    from app.database.base import SessionLocal
    from app.ml.training.versioning import register_model
    from app.models import AudioSegment, Call, EmotionPrediction, Organization, Role, User
    from app.services import config_service
    from app.services.call_service import create_call_from_upload
    from app.workers import pipeline
    from app.workers.pipeline import run_pipeline, run_stage

    import tempfile, pathlib
    tiny = make_tiny_emotion_model(pathlib.Path(tempfile.mkdtemp()) / "tiny")
    db = SessionLocal()
    org = Organization(name="Org Resume")
    db.add(org); db.flush()
    user = User(org_id=org.id, email="r@acme-r.com", hashed_password="x", role=Role.ADMIN.value)
    db.add(user); db.flush()
    m = register_model(db, org_id=org.id, kind="emotion", loader="finetuned", path=str(tiny), base_model=None, dataset_id=None,
                       params={}, labels=["angry", "happy", "neutral", "sad"], language="es", created_by=user.id, status="PRODUCTION")
    config_service.set_section(db, org.id, "audio", {"segmentation": {"batch_size": 4}}, user.id)
    db.commit()
    with open(EXAMPLES / "llamada_ejemplo_estereo.wav", "rb") as f:
        call = create_call_from_upload(db, user, "resume.wav", f, allow_training=False, display_name=None)
    cid = call.id
    for st in ("audio_processing", "diarization", "transcription"):
        assert run_stage(cid, st)
    db.expire_all()
    audio_ts = db.get(Call, cid).checkpoints["audio_processing"]["updated_at"]

    orig = pipeline.iter_window_results

    def flaky(*a, **k):
        for i, batch in enumerate(orig(*a, **k)):
            if i == 2:
                raise RuntimeError("fallo simulado a mitad de la etapa")
            yield batch
    monkeypatch.setattr(pipeline, "iter_window_results", flaky)
    assert run_stage(cid, "emotion_analysis") is False
    db.expire_all()
    c = db.get(Call, cid)
    assert c.status == "ERROR" and c.checkpoints["emotion_analysis"]["status"] == "FAILED"
    done = c.checkpoints["emotion_analysis"]["done"]
    assert done == 8 and c.checkpoints["transcription"]["status"] == "COMPLETED"    # 2 lotes de 4 persistidos
    n_before = db.scalar(select(func.count()).select_from(EmotionPrediction).where(EmotionPrediction.call_id == cid))
    assert 0 < n_before <= 8
    from app.services.serializers import user_error_message
    assert "RuntimeError" not in user_error_message(c) and "fallo simulado" not in user_error_message(c)

    monkeypatch.setattr(pipeline, "iter_window_results", orig)
    assert run_pipeline(cid) is True
    db.expire_all()
    c = db.get(Call, cid)
    assert c.status == "COMPLETED" and c.checkpoints["audio_processing"]["updated_at"] == audio_ts   # audio no se reprocesó
    rows = db.execute(select(EmotionPrediction.speaker_id, EmotionPrediction.start)).all()
    mine = db.execute(select(EmotionPrediction.speaker_id, EmotionPrediction.start).where(EmotionPrediction.call_id == cid)).all()
    assert len(mine) == len(set(mine)) and len(mine) == c.checkpoints["emotion_analysis"]["valid"] > n_before   # sin duplicados
    assert c.model_id == m.id
    db.close()


# ---- una persona vs dos personas (audio mono) ---------------------------------------------------------------
MONO_1 = EXAMPLES / "monologo_ejemplo_una_persona.wav"      # una sola voz (TTS)
MONO_2 = EXAMPLES / "llamada_ejemplo_mono.wav"              # dos voces mezcladas en un canal (TTS)


def analyze(client, headers, path, speakers=None):
    data = {"speakers": speakers} if speakers else {}
    with open(path, "rb") as f:
        r = client.post(f"{API}/calls/upload", headers=headers, files={"file": (path.name, f)}, data=data)
    assert r.status_code == 201, r.text
    cid = r.json()["id"]
    assert wait_completed(client, headers, cid)["status"] == "COMPLETED"
    return cid


def detail(client, headers, cid):
    return client.get(f"{API}/calls/{cid}", headers=headers).json()


@pytest.mark.skipif(not MONO_1.exists(), reason="falta el monólogo de ejemplo")
def test_single_voice_audio_is_analyzed_as_one_person(client, admin_a):
    d = detail(client, admin_a, analyze(client, admin_a, MONO_1))
    assert d["n_speakers"] == 1 and [s["label"] for s in d["speakers"]] == ["SPEAKER_00"]
    assert d["speakers"][0]["name"] == "Persona"                                   # no «Persona 1»
    assert set(d["summary"]["speakers"]) == {"SPEAKER_00"} and d["satisfaction_p1"] is not None and d["satisfaction_p2"] is None
    assert "A nivel de interacción" not in d["summary"]["executive_summary"]
    assert d["summary"]["interaction"]["participants"] == 1 and d["checkpoints"]["diarization"]["n_speakers"] == 1
    assert "diarization" not in d["analysis_quality_detail"]["components"]         # no penaliza ni premia
    cid = d["id"]
    sat = client.get(f"{API}/calls/{cid}/satisfaction", headers=admin_a).json()
    assert set(sat["speakers"]) == {"SPEAKER_00"} and sat["interaction"] is None
    emo = client.get(f"{API}/calls/{cid}/emotions", headers=admin_a).json()["items"]
    assert emo and {e["speaker"] for e in emo} == {"SPEAKER_00"}
    for fmt in ("csv", "xlsx", "json", "pdf"):                                      # exportaciones con un solo participante
        r = client.get(f"{API}/calls/{cid}/export", headers=admin_a, params={"format": fmt})
        assert r.status_code == 200 and len(r.content) > 500, fmt


@pytest.mark.skipif(not MONO_2.exists(), reason="falta la llamada de ejemplo")
def test_two_voice_mono_is_detected_as_two_and_can_be_forced_to_one(client, admin_a):
    cid = analyze(client, admin_a, MONO_2)
    d = detail(client, admin_a, cid)
    assert d["n_speakers"] == 2 and d["diarization_mode"] == "embedding"
    assert d["checkpoints"]["diarization"]["details"]["centroid_similarity"] <= 0.72
    # el usuario indica que es una sola persona -> se re-analiza desde el audio
    r = client.post(f"{API}/calls/{cid}/reanalyze", headers=admin_a, params={"speakers": "1"})
    assert r.status_code == 200, r.text
    assert wait_completed(client, admin_a, cid)["status"] == "COMPLETED"
    d = detail(client, admin_a, cid)
    assert d["n_speakers"] == 1 and d["options"]["speakers"] == "1" and d["satisfaction_p2"] is None
    # y de nuevo a automático -> vuelve a detectar dos
    client.post(f"{API}/calls/{cid}/reanalyze", headers=admin_a, params={"speakers": "auto"})
    assert wait_completed(client, admin_a, cid)["status"] == "COMPLETED"
    assert detail(client, admin_a, cid)["n_speakers"] == 2


@pytest.mark.skipif(not MONO_1.exists(), reason="falta el monólogo de ejemplo")
def test_forcing_two_people_on_a_monologue_is_honored(client, admin_a):
    d = detail(client, admin_a, analyze(client, admin_a, MONO_1, speakers="2"))
    assert d["n_speakers"] == 2 and d["checkpoints"]["diarization"]["details"].get("forced") is True


def test_text_emotion_models_spanish_and_english():
    from app.core.config import get_config_store
    from app.ml.text.emotion_text import TextEmotionAnalyzer
    a = TextEmotionAnalyzer(get_config_store().defaults("emotion"))
    top = lambda r: max(r, key=r.get)
    es = a.analyze_batch(["Estoy muy molesto con el servicio, llevo tres días esperando.", "Muchas gracias, quedó resuelto, excelente atención.",
                          "Buenos días, gracias por comunicarse con atención al cliente.", "ok"], "es")
    assert top(es[0]) == "angry" and top(es[1]) == "happy" and top(es[2]) == "neutral" and es[3] is None   # texto muy corto: sin evidencia
    en = a.analyze_batch(["I am very angry, I have been waiting for three days.", "Thank you so much, this is resolved."], "en")
    assert top(en[0]) == "angry" and top(en[1]) == "happy"
    assert a.analyze_batch(["Bonjour, je voudrais de l'aide"], "fr") == [None]                                # idioma sin modelo: no se inventa
